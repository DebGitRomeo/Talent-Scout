"""Live smoke test (SPEC 10.3). Hits real APIs and spends a little money.

Run with:  pytest -m live        (needs .env with ANTHROPIC_API_KEY and GITHUB_TOKEN, plus
                                  pip install -e ".[embeddings]" for the local embedding model)
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from datetime import UTC, datetime, timedelta

import pandas as pd
import pytest
from dotenv import load_dotenv
from openpyxl import load_workbook

from .conftest import DATA, ROOT

pytestmark = pytest.mark.live

load_dotenv(ROOT / ".env")
CONTRACT = [
    "Source",
    "Date of release",
    "Type",
    "Summary",
    "Level 1 Topic",
    "Level 2 Topic",
    "Level 3 Topic",
    "Public rating/reviews",
]


def _need(var: str) -> None:
    if not os.environ.get(var):
        pytest.skip(f"{var} not set; live smoke test skipped")


def _run(tmp_path, topics, out, rng):
    cmd = [
        sys.executable,
        "-m",
        "topic_scout",
        str(topics),
        str(DATA / "sources.txt"),
        str(out),
        rng,
        "--max-results-per-query",
        "5",
        "--top-k-per-topic",
        "5",
        "--budget-usd",
        "0.50",
        "--config",
        str(ROOT / "config.yaml"),
    ]
    return subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True, timeout=1800, env=os.environ.copy())


def _last_log(out) -> dict:
    df = pd.read_excel(out, sheet_name="Run Log")
    return df.iloc[-1].to_dict()


def test_live_smoke(tmp_path):
    _need("ANTHROPIC_API_KEY")
    _need("GITHUB_TOKEN")
    pytest.importorskip("sentence_transformers")
    end = datetime.now(UTC).date()
    start = end - timedelta(days=30)
    rng = f"{start.isoformat()}:{end.isoformat()}"
    out = tmp_path / "results.xlsx"
    topics = tmp_path / "topics.xlsx"
    shutil.copy(DATA / "topics.xlsx", topics)

    r1 = _run(tmp_path, topics, out, rng)
    assert r1.returncode == 0, r1.stdout + r1.stderr
    wb = load_workbook(out)
    assert {"Results", "Run Log", "Pruned Branches"} <= set(wb.sheetnames)
    df = pd.read_excel(out, sheet_name="Results")
    assert list(df.columns[:8]) == CONTRACT
    assert (df["Platform"].str.contains("github")).any(), "expected at least one GitHub row"
    assert (df["Platform"].str.contains("arxiv")).any(), "expected at least one arXiv row"
    dates = pd.to_datetime(df["Date of release"]).dt.date
    assert all(start <= d <= end for d in dates)
    assert not df.duplicated(["Source", "Level 3 Topic"]).any()
    assert (df["Summary"].astype(str).str.len() > 0).all()
    from topic_scout.io.topics_loader import load_topics

    l3s = load_topics(topics).l3_texts()
    assert set(df["Level 3 Topic"]) <= l3s

    # second identical run: everything cached -> zero new LLM calls, same row count
    r2 = _run(tmp_path, topics, out, rng)
    assert r2.returncode == 0, r2.stdout + r2.stderr
    assert int(_last_log(out)["LLM calls"]) == 0
    assert len(pd.read_excel(out, sheet_name="Results")) == len(df)

    # add one L3 row -> only that node is expanded
    wb_t = load_workbook(topics)
    ws = wb_t["Topics"]
    ws.append(["Agentic AI", "Agent Memory", "Episodic memory retrieval for agents"])
    wb_t.save(topics)
    r3 = _run(tmp_path, topics, out, rng)
    assert r3.returncode == 0, r3.stdout + r3.stderr
    sent = str(_last_log(out)["Expansion nodes sent"])
    assert sent == "Agentic AI > Agent Memory > Episodic memory retrieval for agents"
