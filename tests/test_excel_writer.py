from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from topic_scout.errors import OutputLockedError
from topic_scout.io.excel_writer import CONTRACT_COLUMNS, RESULT_COLUMNS, write_workbook
from topic_scout.models import PrunedBranch, ResultRow

EXPECTED_CONTRACT = [
    "Source",
    "Date of release",
    "Type",
    "Summary",
    "Level 1 Topic",
    "Level 2 Topic",
    "Level 3 Topic",
    "Public rating/reviews",
]


def row(src: str, l3: str = "L3a", summary: str = "s", rel: float = 0.5) -> ResultRow:
    return ResultRow(
        source=src,
        date_of_release=date(2026, 9, 3),
        type="paper",
        summary=summary,
        level1="L1",
        level2="L2",
        level3=l3,
        rating="Not available",
        title="t",
        platform="arxiv.org",
        relevance_score=rel,
    )


def header(ws):
    return [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]


def test_exact_column_order_and_formatting(tmp_path):
    out = tmp_path / "r.xlsx"
    write_workbook(
        out,
        [row("https://a.org/1"), row("https://a.org/2", rel=0.9)],
        {"Run timestamp": "t1"},
        [PrunedBranch("github.com", "L1", "L2", "zero results", "q")],
        run_ts="2026-10-01 10:00:00",
    )
    wb = load_workbook(out)
    assert wb.sheetnames == ["Results", "Run Log", "Pruned Branches"]
    ws = wb["Results"]
    assert CONTRACT_COLUMNS == EXPECTED_CONTRACT
    assert header(ws) == RESULT_COLUMNS
    assert ws["A2"].value == "https://a.org/2"  # higher relevance first within the same L3
    assert ws["A2"].hyperlink.target == "https://a.org/2"
    assert ws["B2"].number_format == "yyyy-mm-dd"
    assert ws["A1"].font.bold and ws["A1"].font.name == "Arial" and ws["A1"].fill.fgColor.rgb.endswith("333333")
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1")
    assert ws["D2"].alignment.wrap_text and ws.column_dimensions["D"].width == 80
    assert [c.value for c in wb["Pruned Branches"]["A"]][1] == "github.com"


def test_upsert_updates_and_appends(tmp_path):
    out = tmp_path / "r.xlsx"
    write_workbook(out, [row("https://a.org/1", summary="old")], {"Run timestamp": "t1"}, [], run_ts="T1")
    s = write_workbook(
        out,
        [row("https://a.org/1", summary="new"), row("https://a.org/9")],
        {"Run timestamp": "t2"},
        [],
        run_ts="T2",
        taxonomy_l3={"L3a"},
    )
    assert (s.new_rows, s.updated_rows, s.total_rows) == (1, 1, 2)
    ws = load_workbook(out)["Results"]
    recs = {r[0]: dict(zip(RESULT_COLUMNS, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert recs["https://a.org/1"]["Summary"] == "new"
    assert recs["https://a.org/1"]["First seen"] == "T1" and recs["https://a.org/1"]["Last updated"] == "T2"
    assert len(list(load_workbook(out)["Run Log"].iter_rows(min_row=2))) == 2


def test_orphaned_rows_counted(tmp_path):
    out = tmp_path / "r.xlsx"
    write_workbook(out, [row("https://a.org/1", l3="Removed topic")], {}, [], run_ts="T1")
    s = write_workbook(out, [row("https://a.org/2")], {}, [], run_ts="T2", taxonomy_l3={"L3a"})
    assert s.orphaned_rows == 1 and s.total_rows == 2


def test_overwrite_mode(tmp_path):
    out = tmp_path / "r.xlsx"
    write_workbook(out, [row("https://a.org/1")], {}, [], run_ts="T1")
    s = write_workbook(out, [row("https://a.org/2")], {}, [], mode="overwrite", run_ts="T2")
    assert s.total_rows == 1
    assert len(list(load_workbook(out)["Run Log"].iter_rows(min_row=2))) == 2  # run log still appends


def test_extra_columns_hidden(tmp_path):
    out = tmp_path / "r.xlsx"
    write_workbook(out, [row("https://a.org/1")], {}, [], extra_columns=False)
    ws = load_workbook(out)["Results"]
    assert ws.column_dimensions["I"].hidden and not ws.column_dimensions["H"].hidden


def test_locked_file_exit_4(tmp_path):
    out = tmp_path / "r.xlsx"
    (tmp_path / "~$r.xlsx").write_text("lock")
    with pytest.raises(OutputLockedError) as e:
        write_workbook(out, [row("https://a.org/1")], {}, [])
    assert e.value.exit_code == 4
