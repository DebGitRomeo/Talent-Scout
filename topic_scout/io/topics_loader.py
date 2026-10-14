"""Load the three-level taxonomy workbook (SPEC 3.1)."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from topic_scout.errors import InputError
from topic_scout.models import Taxonomy

log = logging.getLogger(__name__)

REQUIRED = ("level 1 topic", "level 2 topic", "level 3 topic")


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _header_key(value: Any) -> str:
    return re.sub(r"\s+", " ", _norm(value).lower())


def load_topics(path: Path) -> Taxonomy:
    if not path.exists():
        raise InputError(f"Topics workbook not found: {path}")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types
        raise InputError(f"Cannot open topics workbook {path}: {exc}") from exc
    ws = wb["Topics"] if "Topics" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)

    header_row: tuple[Any, ...] | None = None
    for row in rows_iter:
        if any(_norm(c) for c in row):
            header_row = row
            break
    if header_row is None:
        raise InputError(f"Topics sheet '{ws.title}' in {path} is empty.")
    keys = [_header_key(c) for c in header_row]
    missing = [h for h in REQUIRED if h not in keys]
    if missing:
        raise InputError(
            f"Topics sheet '{ws.title}' is missing header(s): "
            + ", ".join(m.title() for m in missing)
            + ". Expected 'Level 1 Topic', 'Level 2 Topic', 'Level 3 Topic'."
        )
    idx = [keys.index(h) for h in REQUIRED]

    rows: list[tuple[str, str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    last_l1, last_l2 = "", ""
    for raw in rows_iter:
        cells = [_norm(raw[i]) if i < len(raw) else "" for i in idx]
        if not any(cells):
            continue
        l1, l2, l3 = cells
        if l1 and l1 != last_l1:
            last_l2 = ""  # a new L1 must not inherit the previous L1's L2
        l1 = l1 or last_l1
        l2 = l2 or last_l2
        if not l1 or not l2:
            log.warning("Skipping topics row with no Level 1/Level 2 to inherit: %s", cells)
            continue
        last_l1, last_l2 = l1, l2
        l3 = l3 or l2  # blank L3 = search the whole L2 branch
        key = (l1, l2, l3)
        if key in seen:
            log.warning("Dropping duplicate topics row: %s > %s > %s", *key)
            continue
        seen.add(key)
        rows.append(key)
    wb.close()
    if not rows:
        raise InputError(f"No valid topic rows found in {path}.")
    return Taxonomy.from_rows(rows)
