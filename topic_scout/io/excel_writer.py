"""Output workbook (SPEC 7): Results (upsert/overwrite), Run Log (append), Pruned Branches (replace)."""

from __future__ import annotations

import logging
import os
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from topic_scout.errors import OutputLockedError
from topic_scout.models import PrunedBranch, ResultRow

log = logging.getLogger(__name__)

RESULT_COLUMNS: list[str] = [
    "Source",
    "Date of release",
    "Type",
    "Summary",
    "Level 1 Topic",
    "Level 2 Topic",
    "Level 3 Topic",
    "Public rating/reviews",
    "Title",
    "Platform",
    "Relevance score",
    "Rating score",
    "First seen",
    "Last updated",
]
CONTRACT_COLUMNS = RESULT_COLUMNS[:8]
PRUNED_COLUMNS = ["Source", "L1", "L2", "Reason", "Probe query"]
WIDTHS = {
    "Source": 45,
    "Date of release": 14,
    "Type": 11,
    "Summary": 80,
    "Level 1 Topic": 22,
    "Level 2 Topic": 28,
    "Level 3 Topic": 32,
    "Public rating/reviews": 34,
    "Title": 40,
    "Platform": 24,
    "Relevance score": 10,
    "Rating score": 10,
    "First seen": 20,
    "Last updated": 20,
}

FONT = Font(name="Arial", size=10)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="333333")
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")


@dataclass
class WriteSummary:
    total_rows: int
    new_rows: int
    updated_rows: int
    orphaned_rows: int


def _row_to_record(r: ResultRow, run_ts: str) -> dict[str, Any]:
    return {
        "Source": r.source,
        "Date of release": r.date_of_release,
        "Type": r.type,
        "Summary": r.summary,
        "Level 1 Topic": r.level1,
        "Level 2 Topic": r.level2,
        "Level 3 Topic": r.level3,
        "Public rating/reviews": r.rating,
        "Title": r.title,
        "Platform": r.platform,
        "Relevance score": round(r.relevance_score, 4),
        "Rating score": r.rating_score,
        "First seen": run_ts,
        "Last updated": run_ts,
    }


def check_writable(path: Path) -> None:
    lock = path.parent / f"~${path.name}"
    if lock.exists():
        raise OutputLockedError(f"{path} appears to be open in Excel (found {lock.name}). Close it and re-run.")
    if path.exists():
        try:
            with path.open("r+b"):
                pass
        except PermissionError as exc:
            raise OutputLockedError(f"{path} is locked or read-only. Close it (e.g. in Excel) and re-run.") from exc


def _read_sheet(ws: Worksheet) -> list[dict[str, Any]]:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    names = [str(h) if h is not None else "" for h in header]
    out = []
    for values in rows:
        if values is None or all(v is None for v in values):
            continue
        rec = dict(zip(names, values, strict=False))
        out.append(rec)
    return out


def _as_date(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, str):
        try:
            return date.fromisoformat(v[:10])
        except ValueError:
            return v
    return v


def write_workbook(
    path: Path,
    rows: list[ResultRow],
    run_log: dict[str, Any],
    pruned: list[PrunedBranch],
    *,
    mode: str = "upsert",
    extra_columns: bool = True,
    taxonomy_l3: set[str] | None = None,
    run_ts: str | None = None,
) -> WriteSummary:
    run_ts = run_ts or datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    path.parent.mkdir(parents=True, exist_ok=True)
    check_writable(path)

    existing_results: list[dict[str, Any]] = []
    existing_log: list[dict[str, Any]] = []
    if path.exists():
        try:
            wb_old = load_workbook(path)
        except Exception as exc:  # noqa: BLE001
            raise OutputLockedError(f"Cannot read existing output {path}: {exc}") from exc
        if "Results" in wb_old.sheetnames and mode == "upsert":
            existing_results = _read_sheet(wb_old["Results"])
        if "Run Log" in wb_old.sheetnames:
            existing_log = _read_sheet(wb_old["Run Log"])
        wb_old.close()

    merged: dict[tuple[str, str], dict[str, Any]] = {}
    for rec in existing_results:
        rec["Date of release"] = _as_date(rec.get("Date of release"))
        merged[(str(rec.get("Source")), str(rec.get("Level 3 Topic")))] = rec
    new_rows = updated = 0
    for r in rows:
        rec = _row_to_record(r, run_ts)
        key = (rec["Source"], rec["Level 3 Topic"])
        if key in merged:
            old = merged[key]
            for col in ("Summary", "Public rating/reviews", "Rating score", "Relevance score", "Type", "Title"):
                if rec.get(col) not in (None, ""):
                    old[col] = rec[col]
            old["Last updated"] = run_ts
            old.setdefault("First seen", run_ts)
            updated += 1
        else:
            merged[key] = rec
            new_rows += 1

    orphaned = 0
    if taxonomy_l3 is not None:
        orphaned = sum(1 for (_, l3) in merged if l3 not in taxonomy_l3)
    run_log["Orphaned rows"] = orphaned

    def sort_key(rec: dict[str, Any]) -> tuple:
        rel = rec.get("Relevance score")
        return (
            str(rec.get("Level 1 Topic") or ""),
            str(rec.get("Level 2 Topic") or ""),
            str(rec.get("Level 3 Topic") or ""),
            -(float(rel) if isinstance(rel, int | float) else 0.0),
            str(rec.get("Source")),
        )

    result_records = sorted(merged.values(), key=sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    _write_results(ws, result_records, extra_columns)

    log_cols: list[str] = []
    for rec in [*existing_log, run_log]:
        for k in rec:
            if k not in log_cols and k:
                log_cols.append(k)
    _write_table(wb.create_sheet("Run Log"), log_cols, [*existing_log, run_log], wrap_cols=set())
    _write_table(
        wb.create_sheet("Pruned Branches"),
        PRUNED_COLUMNS,
        [
            {"Source": p.source, "L1": p.l1, "L2": p.l2, "Reason": p.reason, "Probe query": p.probe_query}
            for p in pruned
        ],
        wrap_cols=set(),
    )

    fd, tmp_name = tempfile.mkstemp(prefix=".topic_scout_", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as exc:
        raise OutputLockedError(f"Cannot replace {path}; close it (e.g. in Excel) and re-run.") from exc
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)
    return WriteSummary(len(result_records), new_rows, updated, orphaned)


def _style_header(ws: Worksheet, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ncols:
        ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{max(ws.max_row, 1)}"


def _write_results(ws: Worksheet, records: list[dict[str, Any]], extra_columns: bool) -> None:
    ws.append(RESULT_COLUMNS)
    for rec in records:
        ws.append([rec.get(c) for c in RESULT_COLUMNS])
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(RESULT_COLUMNS) + 1):
            ws.cell(row=r, column=c).font = FONT
        src = ws.cell(row=r, column=1)
        if isinstance(src.value, str) and src.value.startswith("http"):
            src.hyperlink = src.value
            src.font = LINK_FONT
        ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        for c in (1, 3, 5, 6, 7, 8, 9, 10):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=c in (8, 9))
    for i, name in enumerate(RESULT_COLUMNS, start=1):
        dim = ws.column_dimensions[get_column_letter(i)]
        dim.width = WIDTHS.get(name, 16)
        if not extra_columns and i > len(CONTRACT_COLUMNS):
            dim.hidden = True
    _style_header(ws, len(RESULT_COLUMNS))


def _write_table(ws: Worksheet, columns: list[str], records: list[dict[str, Any]], wrap_cols: set[str]) -> None:
    ws.append(columns)
    for rec in records:
        ws.append([_cell_value(rec.get(c)) for c in columns])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for i, _ in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 28
    _style_header(ws, len(columns))


def _cell_value(v: Any) -> Any:
    if isinstance(v, list | tuple):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return "; ".join(f"{k}={val}" for k, val in v.items())
    return v
